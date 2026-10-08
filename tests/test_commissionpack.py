import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from commissionpack import FREE_POINT_LIMIT
from commissionpack.generator import ProjectInfo, build_package, select_points
from commissionpack.i18n import LANGS, T
from commissionpack.license import check_license
from commissionpack.model import AI, AO, DI, DO
from commissionpack.parser import parse_address, parse_file

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


@pytest.mark.parametrize("raw,expected", [
    ("%I0.0", ("%I0.0", DI)),
    ("%Q4.7", ("%Q4.7", DO)),
    ("%IW64", ("%IW64", AI)),
    ("%QW80", ("%QW80", AO)),
    ("%E1.2", ("%I1.2", DI)),       # German mnemonic
    ("%A0.1", ("%Q0.1", DO)),
    ("I 3.4", ("%I3.4", DI)),
    ("%I0.0:P", ("%I0.0", DI)),
    ("%PIW256", ("%PIW256", AI)),
    ("%M10.0", None),
    ("%MD20", None),
    ("%IB0", None),
    ("", None),
    ("DB1.DBX0.0", None),
])
def test_parse_address(raw, expected):
    assert parse_address(raw) == expected


def test_parse_tia_export():
    r = parse_file(SAMPLES / "tia_plc_tags_sample.xlsx")
    assert r.source_format == "tia_portal"
    assert (r.count(DI), r.count(DO), r.count(AI), r.count(AO)) == (20, 9, 4, 2)
    assert r.skipped == 3
    estop = next(p for p in r.points if p.tag == "EStop_Panel_CH1")
    assert estop.safety and estop.device == "safety"
    assert next(p for p in r.points if p.tag == "YV1_Cyl1_Extend").device == "valve"


def test_parse_generic_csv_with_ranges():
    r = parse_file(SAMPLES / "io_list_generic_tr.csv")
    assert r.source_format == "generic"
    assert len(r.points) == 35
    tt = next(p for p in r.points if p.tag == "TT101_Temp")
    assert (tt.range_min, tt.range_max, tt.unit, tt.location) == (0.0, 150.0, "°C", "+CP1")


def test_missing_header_raises(tmp_path):
    f = tmp_path / "bad.csv"
    f.write_text("a;b\n1;2\n", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_file(f)


@pytest.mark.parametrize("lang", LANGS)
def test_build_package_licensed(lang):
    r = parse_file(SAMPLES / "tia_plc_tags_sample.xlsx")
    data = build_package(r, ProjectInfo(project="Test"), lang=lang, licensed=True)
    wb = load_workbook(io.BytesIO(data))
    assert len(wb.sheetnames) == 10
    di = wb[T["sh_DI"][lang]]
    assert di["B5"].value == "EStop_Panel_CH1"
    assert di.max_row == 4 + 20
    ai = wb[T["sh_AI"][lang]]
    assert ai.max_row == 4 + 4 * 5  # 5 test points per channel
    assert ai["K9"].value.startswith("=IF(")


def test_free_edition_limit():
    r = parse_file(SAMPLES / "tia_plc_tags_sample.xlsx")
    chosen = select_points(r.points, FREE_POINT_LIMIT)
    assert len(chosen) == FREE_POINT_LIMIT
    assert {p.kind for p in chosen} == {DI, DO, AI, AO}
    wb = load_workbook(io.BytesIO(build_package(r, ProjectInfo(), lang="en", licensed=False)))
    assert "FREE EDITION" in wb["DI Checkout"]["A3"].value


def test_all_texts_have_all_languages():
    for key, entry in T.items():
        assert set(entry) == set(LANGS), key
        if key.startswith("sh_"):
            assert all(len(v) <= 31 for v in entry.values())


def test_license(monkeypatch):
    monkeypatch.delenv("POLAR_ORGANIZATION_ID", raising=False)
    monkeypatch.setenv("CP_DEV_LICENSE_KEYS", "ABC, DEF")
    assert check_license("DEF").valid
    assert not check_license("").valid
    assert not check_license("XYZ").valid
