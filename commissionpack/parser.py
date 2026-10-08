"""Read an IO list.

Supported inputs:
- TIA Portal PLC tag table export (.xlsx, sheet "PLC Tags",
  columns Name / Data Type / Logical Address / Comment).
- Generic IO list (.xlsx / .xlsm / .csv) with flexible column names
  in English, Turkish or German.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from .classify import classify
from .model import AI, AO, DI, DO, IOPoint, ParseResult

# Column name aliases (lowercase, no spaces/underscores/dots)
ALIASES = {
    "tag": ["name", "tag", "tagname", "symbol", "sembol", "etiket", "symbolname", "variable", "değişken", "degisken"],
    "address": ["logicaladdress", "address", "adres", "adresse", "absoluteaddress", "operand", "ioaddress", "plcaddress"],
    "data_type": ["datatype", "type", "veritipi", "tip", "datentyp", "typ"],
    "comment": ["comment", "description", "açıklama", "aciklama", "kommentar", "beschreibung", "desc", "text"],
    "range_min": ["rangemin", "min", "minvalue", "eumin", "scalemin", "aralıkmin", "aralikmin"],
    "range_max": ["rangemax", "max", "maxvalue", "eumax", "scalemax", "aralıkmax", "aralikmax"],
    "unit": ["unit", "birim", "einheit", "eu", "engunit"],
    "location": ["location", "panel", "cabinet", "pano", "konum", "ort", "schrank", "junctionbox"],
}

# %I0.0  %Q4.7  %IW64  %QD10  %E0.0 (German)  %PIW256  I 0.0  %I0.0:P
ADDR_RE = re.compile(
    r"^%?\s*(P?)(I|Q|E|A)\s*(B|W|D)?\s*(\d+)(?:\.(\d))?(?::P)?$",
    re.IGNORECASE,
)


def _norm_header(text: object) -> str:
    s = str(text or "").strip().lower()
    return re.sub(r"[\s_.\-/()]+", "", s)


def parse_address(raw: str) -> tuple[str, str] | None:
    """Return (normalized_address, kind) or None when not a physical IO."""
    s = str(raw or "").strip().replace(" ", "")
    m = ADDR_RE.match(s)
    if not m:
        return None
    periph, area, size, byte, bit = m.groups()
    area = area.upper()
    area = {"E": "I", "A": "Q"}.get(area, area)  # German mnemonics -> IEC
    size = (size or "").upper()
    if size == "":
        if bit is None:
            return None
        kind = DI if area == "I" else DO
        return f"%{area}{byte}.{bit}", kind
    if size == "B":
        return None  # byte access: not a single test point
    if bit is not None:
        return None
    kind = AI if area == "I" else AO
    prefix = "P" if periph else ""
    return f"%{prefix}{area}{size}{byte}", kind


def _to_float(v: object) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def _map_columns(headers: list[object]) -> dict[str, int]:
    normed = [_norm_header(h) for h in headers]
    mapping: dict[str, int] = {}
    for key, names in ALIASES.items():
        targets = {_norm_header(n) for n in names}
        for idx, h in enumerate(normed):
            if h in targets and idx not in mapping.values():
                mapping[key] = idx
                break
    return mapping


def _rows_to_points(rows: list[list[object]], result: ParseResult) -> None:
    # Find the header row in the first 10 rows
    header_idx = None
    mapping: dict[str, int] = {}
    for i, row in enumerate(rows[:10]):
        m = _map_columns(list(row))
        if "address" in m and "tag" in m:
            header_idx, mapping = i, m
            break
    if header_idx is None:
        raise ValueError(
            "Could not find the header row. The file needs at least a tag/name column "
            "and an address column (e.g. 'Name' and 'Logical Address')."
        )

    def cell(row: list[object], key: str) -> object:
        idx = mapping.get(key)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    seen: set[str] = set()
    for row in rows[header_idx + 1:]:
        if not any(c not in (None, "") for c in row):
            continue
        tag = str(cell(row, "tag") or "").strip()
        parsed = parse_address(str(cell(row, "address") or ""))
        if not tag or parsed is None:
            result.skipped += 1
            continue
        address, kind = parsed
        if address in seen:
            result.warnings.append(f"Duplicate address {address} (tag '{tag}') - kept both rows.")
        seen.add(address)
        point = IOPoint(
            tag=tag,
            address=address,
            kind=kind,
            data_type=str(cell(row, "data_type") or "").strip(),
            comment=str(cell(row, "comment") or "").strip(),
            range_min=_to_float(cell(row, "range_min")),
            range_max=_to_float(cell(row, "range_max")),
            unit=str(cell(row, "unit") or "").strip(),
            location=str(cell(row, "location") or "").strip(),
        )
        point.device, point.safety = classify(point)
        result.points.append(point)


def _read_csv(data: bytes) -> list[list[object]]:
    text = None
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    assert text is not None
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def parse_file(source: str | Path | BinaryIO, filename: str | None = None) -> ParseResult:
    """Parse a file path or a file-like object (e.g. a Streamlit upload)."""
    if isinstance(source, (str, Path)):
        filename = filename or str(source)
        data = Path(source).read_bytes()
    else:
        data = source.read()
        filename = filename or getattr(source, "name", "upload.xlsx")

    result = ParseResult()
    suffix = Path(filename).suffix.lower()

    if suffix == ".csv":
        result.source_format = "generic"
        _rows_to_points(_read_csv(data), result)
    elif suffix in (".xlsx", ".xlsm"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        target = next((n for n in sheet_names if n.strip().lower() == "plc tags"), None)
        result.source_format = "tia_portal" if target else "generic"
        ws = wb[target] if target else wb[sheet_names[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        _rows_to_points(rows, result)
    else:
        raise ValueError(f"Unsupported file type '{suffix}'. Use .xlsx, .xlsm or .csv.")

    if not result.points:
        result.warnings.append("No physical IO points (%I, %Q, %IW, %QW ...) were found.")
    result.points.sort(key=lambda p: (("DI", "DO", "AI", "AO").index(p.kind), _addr_key(p.address)))
    return result


def _addr_key(address: str) -> tuple[int, int]:
    m = re.search(r"(\d+)(?:\.(\d))?$", address)
    if not m:
        return (0, 0)
    return (int(m.group(1)), int(m.group(2) or 0))
