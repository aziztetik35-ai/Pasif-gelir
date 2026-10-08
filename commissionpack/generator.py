"""Build the commissioning workbook (.xlsx) from parsed IO points."""

from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from . import FREE_POINT_LIMIT, __version__
from .i18n import EXPECTED, FAT_ITEMS, IO_REF_TEXT, SAT_ITEMS, device_name, t, test_text
from .model import AI, AO, DI, DO, KINDS, IOPoint, ParseResult

# ---------------------------------------------------------------- styles
DARK = "1F3864"
F_HEADER = PatternFill("solid", fgColor=DARK)
F_INPUT = PatternFill("solid", fgColor="FFF9DB")
F_SAFETY = PatternFill("solid", fgColor="F8D7DA")
F_SECTION = PatternFill("solid", fgColor="D9E1F2")
F_FREE = PatternFill("solid", fgColor="FFE699")
FONT_HEADER = Font(bold=True, color="FFFFFF")
FONT_TITLE = Font(bold=True, size=14, color=DARK)
FONT_BIG = Font(bold=True, size=20, color=DARK)
FONT_BOLD = Font(bold=True)
FONT_SMALL = Font(size=9, color="595959")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

HEADER_ROW = 4
FIRST_ROW = 5
ANALOG_POINTS = (0.0, 0.25, 0.5, 0.75, 1.0)
SIGNALS = ("4-20 mA", "0-20 mA", "0-10 V", "+/-10 V")


@dataclass
class ProjectInfo:
    project: str = ""
    customer: str = ""
    machine: str = ""
    plc: str = ""
    doc_no: str = ""
    revision: str = "0"
    date: str = field(default_factory=lambda: dt.date.today().isoformat())
    prepared_by: str = ""


def _q(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'"


def select_points(points: list[IOPoint], limit: int) -> list[IOPoint]:
    """Pick up to `limit` points, spread over the signal kinds (free edition)."""
    if len(points) <= limit:
        return list(points)
    # Round-robin over the kinds, so small groups (AI/AO) are shown too.
    queues = [[p for p in points if p.kind == k] for k in KINDS]
    chosen: list[IOPoint] = []
    while len(chosen) < limit:
        for q in queues:
            if q and len(chosen) < limit:
                chosen.append(q.pop(0))
    # keep the original order
    ids = {id(p) for p in chosen}
    return [p for p in points if id(p) in ids]


class PackageBuilder:
    def __init__(self, result: ParseResult, info: ProjectInfo, lang: str = "en",
                 licensed: bool = False, source_name: str = ""):
        if lang not in ("en", "tr", "de"):
            raise ValueError("lang must be en, tr or de")
        self.lang = lang
        self.info = info
        self.licensed = licensed
        self.source_name = source_name
        self.all_points = result.points
        self.points = result.points if licensed else select_points(result.points, FREE_POINT_LIMIT)
        self.wb = Workbook()
        self.names = {k: t(f"sh_{k}", lang) for k in KINDS}
        # Summary rows are fixed: one row per signal kind, in KINDS order.
        self.summary_rows = {k: FIRST_ROW + i for i, k in enumerate(KINDS)}
        self.ranges: dict[str, tuple[int, int]] = {}  # kind -> (first_row, last_row)

    # ------------------------------------------------------------ helpers
    def _tx(self, key: str) -> str:
        return t(key, self.lang)

    def _sheet_header(self, ws: Worksheet, title: str, ncols: int) -> None:
        ws["A1"] = f"{title} - {self.info.project}".rstrip(" -")
        ws["A1"].font = FONT_TITLE
        ws["A2"] = (f"{self._tx('customer')}: {self.info.customer}   |   {self._tx('machine')}: {self.info.machine}"
                    f"   |   {self._tx('doc_no')}: {self.info.doc_no}   |   {self._tx('revision')}: {self.info.revision}")
        ws["A2"].font = FONT_SMALL
        if not self.licensed:
            ws["A3"] = t("free_notice", self.lang).format(n=FREE_POINT_LIMIT)
            ws["A3"].font = FONT_BOLD
            for c in range(1, ncols + 1):
                ws.cell(row=3, column=c).fill = F_FREE

    def _table_header(self, ws: Worksheet, headers: list[str], widths: list[int]) -> None:
        for i, (h, w) in enumerate(zip(headers, widths), start=1):
            c = ws.cell(row=HEADER_ROW, column=i, value=h)
            c.fill, c.font, c.alignment, c.border = F_HEADER, FONT_HEADER, CENTER, BORDER
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=FIRST_ROW, column=3)
        ws.row_dimensions[HEADER_ROW].height = 32

    def _print_setup(self, ws: Worksheet, last_col: int, last_row: int) -> None:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
        ws.print_area = f"A1:{get_column_letter(last_col)}{max(last_row, HEADER_ROW)}"
        ws.oddFooter.left.text = f"{self.info.project} {self.info.doc_no}".strip()
        ws.oddFooter.center.text = "&P / &N"
        ws.oddFooter.right.text = f"CommissionPack {__version__}"

    def _status_rules(self, ws: Worksheet, rng: str) -> None:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'],
                                                      fill=PatternFill("solid", fgColor="C6EFCE")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NOK"'],
                                                      fill=PatternFill("solid", fgColor="FFC7CE")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"N/A"'],
                                                      fill=PatternFill("solid", fgColor="E7E6E6")))

    @staticmethod
    def _status_validation(ws: Worksheet, rng: str) -> None:
        dv = DataValidation(type="list", formula1='"OK,NOK,N/A"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng)

    # ------------------------------------------------------------ sheets
    def _digital_sheet(self, kind: str) -> None:
        name = self.names[kind]
        ws = self.wb.create_sheet(name)
        keys = ["no", "tag", "address", "comment", "location", "device", "test", "expected",
                "status", "tested_by", "date", "remarks"]
        widths = [5, 24, 10, 32, 12, 18, 58, 30, 9, 14, 12, 26]
        self._sheet_header(ws, name, len(keys))
        self._table_header(ws, [self._tx(k) for k in keys], widths)
        pts = [p for p in self.points if p.kind == kind]
        row = FIRST_ROW
        for n, p in enumerate(pts, start=1):
            values = [n, p.tag, p.address, p.comment, p.location, device_name(p.device, self.lang),
                      test_text(kind, p.device, self.lang), EXPECTED[kind][self.lang], None, None, None, None]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                c.border, c.alignment = BORDER, WRAP
                if col >= 9:
                    c.fill = F_INPUT
                elif p.safety:
                    c.fill = F_SAFETY
            row += 1
        last = max(row - 1, FIRST_ROW)
        self.ranges[kind] = (FIRST_ROW, last)
        self._status_validation(ws, f"I{FIRST_ROW}:I{last}")
        self._status_rules(ws, f"I{FIRST_ROW}:I{last}")
        ws.auto_filter.ref = f"A{HEADER_ROW}:L{last}"
        self._print_setup(ws, len(keys), last)

    def _analog_sheet(self, kind: str) -> None:
        name = self.names[kind]
        ws = self.wb.create_sheet(name)
        keys = ["no", "tag", "address", "comment", "signal", "rmin", "rmax", "unit", "point", "exp_signal",
                "raw", "exp_eu", "measured", "deviation", "tolerance", "status", "tested_by", "date", "remarks"]
        widths = [5, 22, 10, 28, 10, 9, 9, 7, 8, 10, 10, 10, 10, 10, 9, 8, 12, 11, 22]
        self._sheet_header(ws, name, len(keys))
        method = "test_ai" if kind == AI else "test_ao"
        ws["A3" if self.licensed else "H3"] = t(method, self.lang)
        ws["A3" if self.licensed else "H3"].font = FONT_SMALL
        self._table_header(ws, [self._tx(k) for k in keys], widths)
        pts = [p for p in self.points if p.kind == kind]
        row = FIRST_ROW
        sig_dv = DataValidation(type="list", formula1='"' + ",".join(SIGNALS) + '"', allow_blank=False)
        ws.add_data_validation(sig_dv)
        for n, p in enumerate(pts, start=1):
            first = row
            for i, pct in enumerate(ANALOG_POINTS):
                r = row
                is_first = i == 0
                cells = {
                    1: n if is_first else None,
                    2: p.tag,
                    3: p.address,
                    4: p.comment if is_first else None,
                    5: "4-20 mA" if is_first else f"=E{first}",
                    6: (p.range_min if p.range_min is not None else 0) if is_first else f"=F{first}",
                    7: (p.range_max if p.range_max is not None else 100) if is_first else f"=G{first}",
                    8: p.unit if is_first else f"=H{first}",
                    9: pct,
                    10: f'=IF(E{r}="0-10 V",10*I{r},IF(E{r}="0-20 mA",20*I{r},IF(E{r}="+/-10 V",-10+20*I{r},4+16*I{r})))',
                    11: f'=IF(E{r}="+/-10 V",ROUND(-27648+55296*I{r},0),ROUND(27648*I{r},0))',
                    12: f"=F{r}+I{r}*(G{r}-F{r})",
                    13: None,
                    14: f'=IF(OR(M{r}="",G{r}=F{r}),"",ABS(M{r}-L{r})/ABS(G{r}-F{r})*100)',
                    15: 0.5 if is_first else f"=O{first}",
                    16: f'=IF(N{r}="","",IF(N{r}<=O{r},"OK","NOK"))',
                    17: None, 18: None, 19: None,
                }
                for col, v in cells.items():
                    c = ws.cell(row=r, column=col, value=v)
                    c.border, c.alignment = BORDER, WRAP
                    if col in (13, 17, 18, 19) or (is_first and col in (5, 6, 7, 8, 15)):
                        c.fill = F_INPUT
                    elif p.safety:
                        c.fill = F_SAFETY
                    if is_first:
                        c.border = Border(left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color=DARK))
                ws.cell(row=r, column=9).number_format = "0%"
                for col in (10, 12, 13, 14):
                    ws.cell(row=r, column=col).number_format = "0.00"
                row += 1
            sig_dv.add(f"E{first}")
        last = max(row - 1, FIRST_ROW)
        self.ranges[kind] = (FIRST_ROW, last)
        self._status_validation(ws, f"P{FIRST_ROW}:P{last}")
        self._status_rules(ws, f"P{FIRST_ROW}:P{last}")
        ws.auto_filter.ref = f"A{HEADER_ROW}:S{last}"
        self._print_setup(ws, len(keys), last)

    def _summary_sheet(self, ws: Worksheet) -> None:
        self._sheet_header(ws, self._tx("sh_summary"), 7)
        labels = [self._tx("section"), self._tx("total"), "OK", "NOK", "N/A", self._tx("open"), self._tx("complete")]
        self._table_header(ws, labels, [34, 10, 10, 10, 10, 10, 14])
        ws.freeze_panes = None
        row = FIRST_ROW
        for kind in KINDS:
            first, last = self.ranges[kind]
            sh = _q(self.names[kind])
            status_col = "I" if kind in (DI, DO) else "P"
            st = f"{sh}!{status_col}{first}:{status_col}{last}"
            values = [
                IO_REF_TEXT[kind][self.lang],
                f"=COUNTA({sh}!B{first}:B{last})",
                f'=COUNTIF({st},"OK")',
                f'=COUNTIF({st},"NOK")',
                f'=COUNTIF({st},"N/A")',
                f"=B{row}-C{row}-D{row}-E{row}",
                f'=IF(B{row}=0,"",(C{row}+E{row})/B{row})',
            ]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                c.border = BORDER
            ws.cell(row=row, column=7).number_format = "0%"
            row += 1
        # total row
        ws.cell(row=row, column=1, value=self._tx("total")).font = FONT_BOLD
        for col in range(2, 7):
            L = get_column_letter(col)
            ws.cell(row=row, column=col, value=f"=SUM({L}{FIRST_ROW}:{L}{row - 1})").font = FONT_BOLD
        ws.cell(row=row, column=7, value=f'=IF(B{row}=0,"",(C{row}+E{row})/B{row})').number_format = "0%"
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = BORDER
            ws.cell(row=row, column=col).fill = F_SECTION
        ws.conditional_formatting.add(f"G{FIRST_ROW}:G{row}",
                                      DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                                  color="63BE7B"))
        row += 2
        ws.cell(row=row, column=1, value=self._tx("points_note")).font = FONT_SMALL
        row += 1
        punch = _q(self._tx("sh_punch"))
        ws.cell(row=row, column=1, value=f"{self._tx('open')} A ({self._tx('sh_punch')})").font = FONT_BOLD
        ws.cell(row=row, column=2, value=f'=COUNTIFS({punch}!F{FIRST_ROW + 1}:F500,"A",{punch}!I{FIRST_ROW + 1}:I500,"<>CLOSED")')
        self._print_setup(ws, 7, row)

    def _protocol_sheet(self, key: str, items: list[dict[str, object]]) -> None:
        name = self._tx(key)
        ws = self.wb.create_sheet(name)
        keys = ["no", "section", "item", "criterion", "status", "remarks", "tested_by", "date"]
        widths = [5, 18, 60, 36, 9, 30, 14, 12]
        self._sheet_header(ws, name, len(keys))
        self._table_header(ws, [self._tx(k) for k in keys], widths)
        ws.freeze_panes = ws.cell(row=FIRST_ROW, column=1)
        summary = _q(self._tx("sh_summary"))
        n_safety = sum(1 for p in self.all_points if p.safety)
        row = FIRST_ROW
        prev_section = None
        for n, it in enumerate(items, start=1):
            if "io" in it:
                kind = str(it["io"])
                sr = self.summary_rows[kind]
                count = sum(1 for p in self.points if p.kind == kind)
                section = IO_REF_TEXT["section"][self.lang]
                item = f"{IO_REF_TEXT[kind][self.lang]} ({count} {self._tx('io_points')})"
                crit = IO_REF_TEXT["crit"][self.lang].format(sheet=self.names[kind])
                status = f'=IF({summary}!B{sr}=0,"N/A",IF({summary}!G{sr}>=1,"OK",""))'
                remarks = f'=IF({summary}!B{sr}=0,"","{self._tx("complete")}: "&TEXT({summary}!G{sr},"0%"))'
            else:
                section = it["s"][self.lang]  # type: ignore[index]
                item = it["i"][self.lang]  # type: ignore[index]
                crit = it["c"][self.lang]  # type: ignore[index]
                status, remarks = None, None
                if "safety function" in it["i"]["en"] and n_safety:  # type: ignore[index]
                    remarks = f"{self._tx('safety_points')}: {n_safety}"
            values = [n, section, item, crit, status, remarks, None, None]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                c.border, c.alignment = BORDER, WRAP
                if col >= 5 and v is None:
                    c.fill = F_INPUT
                if section != prev_section and col <= 2:
                    c.fill = F_SECTION
            prev_section = section
            row += 1
        last = row - 1
        self._status_validation(ws, f"E{FIRST_ROW}:E{last}")
        self._status_rules(ws, f"E{FIRST_ROW}:E{last}")
        self._print_setup(ws, len(keys), last)

    def _punch_sheet(self) -> None:
        name = self._tx("sh_punch")
        ws = self.wb.create_sheet(name)
        keys = ["no", "ref", "description", "location", "responsible", "category", "due", "remarks", "status",
                "closed_date"]
        widths = [5, 16, 60, 14, 16, 12, 12, 26, 10, 12]
        self._sheet_header(ws, name, len(keys))
        self._table_header(ws, [self._tx(k) for k in keys], widths)
        ws.cell(row=FIRST_ROW, column=1, value=self._tx("punch_note")).font = FONT_SMALL
        first = FIRST_ROW + 1
        last = first + 59
        for r in range(first, last + 1):
            for col in range(1, len(keys) + 1):
                c = ws.cell(row=r, column=col, value=(r - first + 1) if col == 1 else None)
                c.border, c.alignment = BORDER, WRAP
                if col > 1:
                    c.fill = F_INPUT
        cat = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
        st = DataValidation(type="list", formula1='"OPEN,CLOSED"', allow_blank=True)
        ws.add_data_validation(cat)
        ws.add_data_validation(st)
        cat.add(f"F{first}:F{last}")
        st.add(f"I{first}:I{last}")
        ws.conditional_formatting.add(f"I{first}:I{last}", CellIsRule(operator="equal", formula=['"OPEN"'],
                                                                      fill=PatternFill("solid", fgColor="FFC7CE")))
        ws.conditional_formatting.add(f"I{first}:I{last}", CellIsRule(operator="equal", formula=['"CLOSED"'],
                                                                      fill=PatternFill("solid", fgColor="C6EFCE")))
        self._print_setup(ws, len(keys), last)

    def _signoff_sheet(self) -> None:
        name = self._tx("sh_signoff")
        ws = self.wb.create_sheet(name)
        self._sheet_header(ws, name, 5)
        for col, w in enumerate([28, 24, 24, 26, 14], start=1):
            ws.column_dimensions[get_column_letter(col)].width = w
        row = HEADER_ROW
        for block in ("sh_fat", "sh_sat"):
            ws.cell(row=row, column=1, value=self._tx(block)).font = FONT_TITLE
            row += 1
            ws.cell(row=row, column=1, value=f"{self._tx('decision')}:").font = FONT_BOLD
            ws.cell(row=row, column=2, value=self._tx("decision_opts"))
            row += 1
            for col, k in enumerate(["role", "company", "name", "signature", "date"], start=1):
                c = ws.cell(row=row, column=col, value=self._tx(k))
                c.fill, c.font, c.border, c.alignment = F_HEADER, FONT_HEADER, BORDER, CENTER
            row += 1
            for role in ("supplier", "customer_rep", "third_party"):
                for col in range(1, 6):
                    c = ws.cell(row=row, column=col, value=self._tx(role) if col == 1 else None)
                    c.border = BORDER
                    if col > 1:
                        c.fill = F_INPUT
                ws.row_dimensions[row].height = 36
                row += 1
            row += 2
        self._print_setup(ws, 5, row)

    def _cover_sheet(self, ws: Worksheet) -> None:
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 60
        ws["A1"] = self._tx("doc_title")
        ws["A1"].font = FONT_BIG
        ws.row_dimensions[1].height = 34
        row = 3
        fields = [
            ("project", self.info.project), ("customer", self.info.customer), ("machine", self.info.machine),
            ("plc", self.info.plc), ("doc_no", self.info.doc_no), ("revision", self.info.revision),
            ("date", self.info.date), ("prepared_by", self.info.prepared_by),
            ("source_file", self.source_name),
        ]
        for key, val in fields:
            a = ws.cell(row=row, column=1, value=self._tx(key))
            b = ws.cell(row=row, column=2, value=val)
            a.font, a.fill = FONT_BOLD, F_SECTION
            a.border = b.border = BORDER
            row += 1
        row += 1
        ws.cell(row=row, column=1, value=self._tx("contents")).font = FONT_TITLE
        row += 1
        for kind in KINDS:
            count = sum(1 for p in self.points if p.kind == kind)
            ws.cell(row=row, column=1, value=self.names[kind]).border = BORDER
            ws.cell(row=row, column=2, value=f"{count} {self._tx('io_points')}").border = BORDER
            row += 1
        for key in ("sh_fat", "sh_sat", "sh_punch", "sh_signoff", "sh_summary"):
            ws.cell(row=row, column=1, value=self._tx(key)).border = BORDER
            ws.cell(row=row, column=2).border = BORDER
            row += 1
        row += 1
        n_safety = sum(1 for p in self.points if p.safety)
        ws.cell(row=row, column=1, value=self._tx("safety_points")).font = FONT_BOLD
        ws.cell(row=row, column=2, value=n_safety).fill = F_SAFETY
        row += 2
        if not self.licensed:
            c = ws.cell(row=row, column=1, value=t("free_notice", self.lang).format(n=FREE_POINT_LIMIT)
                        + f"  ({len(self.points)} / {len(self.all_points)})")
            c.font, c.fill = FONT_BOLD, F_FREE
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            row += 2
        c = ws.cell(row=row, column=1, value=self._tx("disclaimer"))
        c.alignment, c.font = WRAP, FONT_SMALL
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = 45
        row += 2
        ws.cell(row=row, column=1, value=f"CommissionPack {__version__} - {dt.datetime.now():%Y-%m-%d %H:%M}").font = FONT_SMALL
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ------------------------------------------------------------ build
    def build(self) -> Workbook:
        cover = self.wb.active
        cover.title = self._tx("sh_cover")
        summary = self.wb.create_sheet(self._tx("sh_summary"))
        for kind in (DI, DO):
            self._digital_sheet(kind)
        for kind in (AI, AO):
            self._analog_sheet(kind)
        self._protocol_sheet("sh_fat", FAT_ITEMS)
        self._protocol_sheet("sh_sat", SAT_ITEMS)
        self._punch_sheet()
        self._signoff_sheet()
        # The summary needs the row ranges of the IO sheets, so build it last.
        self._summary_sheet(summary)
        self._cover_sheet(cover)
        self.wb.properties.title = self._tx("doc_title")
        self.wb.properties.creator = f"CommissionPack {__version__}"
        return self.wb


def build_package(result: ParseResult, info: ProjectInfo, lang: str = "en", licensed: bool = False,
                  source_name: str = "") -> bytes:
    """Return the finished .xlsx file as bytes."""
    wb = PackageBuilder(result, info, lang=lang, licensed=licensed, source_name=source_name).build()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
